#!/usr/bin/env python3
"""Bundle the 2026 candidate CSVs into one Excel workbook for easy browsing.

    python3 scripts/elections/build_workbook_2026.py

Reads data/elections/2026/*.csv and indiana/*.csv; writes
data/elections/2026/candidates-2026.xlsx.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data" / "elections" / "2026"

SHEETS = [
    ("US Senate", DATA / "us-senate.csv", None),
    ("US House", DATA / "us-house.csv", None),
    ("Indiana - Federal & State", DATA / "indiana" / "candidates.csv",
     lambda r: r["level"] in ("federal", "statewide", "state legislative")),
    ("Indiana - All Offices", DATA / "indiana" / "candidates.csv", None),
    ("Indiana - Races", DATA / "indiana" / "offices.csv", None),
]


def main():
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1A2035")
    for title, path, keep in SHEETS:
        with path.open(encoding="utf-8") as f:
            rows = list(csv.DictReader(f))
        rows = [r for r in rows if keep is None or keep(r)]
        ws = wb.create_sheet(title)
        cols = list(rows[0].keys())
        ws.append(cols)
        for r in rows:
            ws.append([int(r[c]) if c in ("seats", "candidates", "write_ins") and r[c].isdigit() else r[c] for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, c in enumerate(cols, 1):
            width = max([len(c)] + [len(str(r[c])) for r in rows[:500]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
    out = DATA / "candidates-2026.xlsx"
    wb.save(out)
    print(f"wrote {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
